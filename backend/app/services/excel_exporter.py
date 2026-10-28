"""Excel (.xlsx) export with openpyxl.

Every cell goes through ``sanitize_cell`` (formula-injection protection). Sheets get a styled, frozen
header row, auto-filter and readable column widths."""
from __future__ import annotations

import io
import json
from collections import Counter, defaultdict
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.job import Job, JobUrl
from app.models.lead import Lead
from app.utils.security import sanitize_cell

NAVY = "0F2A4A"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14, color=NAVY)

LEAD_COLUMNS = [
    ("Business Name", lambda l: l.business_name),
    ("Email Address", lambda l: l.email),
    ("Original Email (as found)", lambda l: l.original_email),
    ("Email Domain", lambda l: l.email_domain),
    ("Role-Based", lambda l: "Yes" if l.is_role_based else "No"),
    ("Website", lambda l: l.website),
    ("Contact Page URL", lambda l: l.contact_page_url),
    ("Business Category", lambda l: l.category),
    ("City", lambda l: l.city),
    ("State", lambda l: l.state),
    ("Country", lambda l: l.country),
    ("Address", lambda l: l.address),
    ("Phone", lambda l: l.phone),
    ("Source Type", lambda l: l.source_type),
    ("Source Filename", lambda l: l.source_filename),
    ("Source URL", lambda l: l.source_url),
    ("Collection Date (UTC)", lambda l: l.collected_at.strftime("%Y-%m-%d %H:%M") if l.collected_at else ""),
    ("Verification Status", lambda l: l.verification_status),
    ("Verification Detail", lambda l: l.verification_detail),
    ("Tags", lambda l: l.tags),
    ("Notes", lambda l: l.notes),
    ("Test Data", lambda l: "YES - sample data" if l.is_test else ""),
]


def _write_sheet(ws: Worksheet, headers: Sequence[str], rows: Iterable[Sequence[Any]], widths: Optional[Dict[str, int]] = None) -> int:
    ws.append([sanitize_cell(h) for h in headers])
    count = 0
    for row in rows:
        ws.append([sanitize_cell(v) for v in row])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # belt and braces: never let a string be stored as a formula
        count += 1
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(max(len(headers), 1))}{max(count + 1, 1)}"
    for idx, h in enumerate(headers, start=1):
        longest = len(str(h))
        for r in ws.iter_rows(min_row=2, max_row=min(count + 1, 300), min_col=idx, max_col=idx, values_only=True):
            v = r[0]
            if v is not None:
                longest = max(longest, min(len(str(v)), 70))
        ws.column_dimensions[get_column_letter(idx)].width = min(max(12, longest + 2), 60)
    ws.row_dimensions[1].height = 22
    return count


def _source_summary(lead: Lead) -> str:
    parts = []
    for s in lead.sources[:15]:
        where = s.source_filename or s.source_url or ""
        if s.sheet_name:
            where += f" [{s.sheet_name}" + (f" row {s.row_number}]" if s.row_number else "]")
        elif s.page_number:
            where += f" [page {s.page_number}]"
        elif s.row_number:
            where += f" [row {s.row_number}]"
        if where:
            parts.append(where)
    more = len(lead.sources) - 15
    return "; ".join(parts) + (f"; …and {more} more" if more > 0 else "")


def _original_row(lead: Lead) -> str:
    for s in lead.sources:
        if s.extra:
            return json.dumps(s.extra, ensure_ascii=False)[:3000]
    return ""


def build_workbook(
    leads: Sequence[Lead],
    *,
    scope: str = "leads",
    failed_urls: Optional[Sequence[JobUrl]] = None,
    jobs: Optional[Sequence[Job]] = None,
    title: str = "PackLead Scraper export",
    description: str = "",
    include_extra_sheets: bool = True,
) -> Workbook:
    """scope: 'leads' (Leads [+ Unique Emails + Failed URLs + Summary]) | 'unique_emails' | 'by_domain'."""
    wb = Workbook()
    wb.remove(wb.active)

    if scope == "leads":
        ws = wb.create_sheet("Leads")
        headers = [c[0] for c in LEAD_COLUMNS] + ["All Sources", "Sources Count", "Original Row Data"]
        _write_sheet(
            ws,
            headers,
            ([fn(l) for _, fn in LEAD_COLUMNS] + [_source_summary(l), len(l.sources), _original_row(l)] for l in leads),
        )
    if scope in ("unique_emails",) or (scope == "leads" and include_extra_sheets):
        ws = wb.create_sheet("Unique Emails")
        seen: set = set()
        rows = []
        for l in leads:
            if l.email in seen:
                continue
            seen.add(l.email)
            rows.append([l.email, l.original_email, l.email_domain, "Yes" if l.is_role_based else "No", l.business_name, l.verification_status, l.source_type, len(l.sources)])
        _write_sheet(ws, ["Email Address", "Original Email (as found)", "Email Domain", "Role-Based", "Business Name", "Verification Status", "Source Type", "Times Seen"], rows)
    if scope == "by_domain":
        ws = wb.create_sheet("Emails by Domain")
        groups: Dict[str, List[Lead]] = defaultdict(list)
        for l in leads:
            groups[l.email_domain].append(l)
        rows = []
        for domain in sorted(groups):
            for l in sorted(groups[domain], key=lambda x: x.email):
                rows.append([domain, len(groups[domain]), l.email, l.business_name, l.verification_status, l.website])
        _write_sheet(ws, ["Email Domain", "Emails On Domain", "Email Address", "Business Name", "Verification Status", "Website"], rows)

    if scope == "leads" and include_extra_sheets:
        ws = wb.create_sheet("Failed URLs")
        _write_sheet(
            ws,
            ["URL", "Status", "Reason / Note", "Job ID", "Source Reference", "Pages Visited"],
            ([f.url, f.status, f.error or f.note or "", f.job_id, f.row_ref or "", f.pages_visited] for f in (failed_urls or [])),
        )
        ws = wb.create_sheet("Scraping Summary")
        summary: List[List[Any]] = [
            ["Export", title],
            ["Generated (UTC)", datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")],
            ["Scope", description or "All leads"],
            ["Leads in export", len(leads)],
            ["Distinct email domains", len({l.email_domain for l in leads})],
            ["Role-based addresses", sum(1 for l in leads if l.is_role_based)],
            ["Sample/test records included", sum(1 for l in leads if l.is_test)],
            ["Verification note", "Statuses describe syntax/DNS checks only. No mailbox was contacted; deliverability is NOT confirmed."],
        ]
        for status, n in Counter(l.verification_status for l in leads).most_common():
            summary.append([f"Status: {status}", n])
        for src, n in Counter(l.source_type for l in leads).most_common():
            summary.append([f"Source: {src}", n])
        for job in jobs or []:
            summary.append([f"Job #{job.id} ({job.type})", f"{job.name} - {job.status}; URLs {job.processed_urls}/{job.total_urls}, failed {job.failed_urls}, emails found {job.emails_found}, new leads {job.unique_leads}, duplicates {job.duplicates_removed}"])
        _write_sheet(ws, ["Metric", "Value"], summary)
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 90
    return wb


def workbook_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

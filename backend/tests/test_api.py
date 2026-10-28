import asyncio
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.services.email_validator import EmailValidator
from app.workers.runner import JobRunner


async def fake_mx(domain):
    return "mx", [f"mail.{domain}"], ""


@pytest.fixture
def client(monkeypatch):
    monkeypatch.setattr("app.api.leads.EmailValidator", lambda: EmailValidator(lookup=fake_mx))
    with TestClient(app) as c:
        yield c


def run_worker():
    async def go():
        r = JobRunner(mx_lookup=fake_mx)
        while await r.run_next():
            pass

    asyncio.run(go())


def crawl_fixture(client, urls=("http://127.0.0.1:8099/", "http://localhost:8099/"), **extra):
    res = client.post("/api/jobs/website", json={"urls": list(urls), "delay": 0, "max_pages": 10, **extra})
    assert res.status_code == 201, res.text
    run_worker()
    return res.json()["job"]["id"]


def test_health(client):
    for path in ("/health", "/api/health"):
        r = client.get(path)
        assert r.status_code == 200 and r.json()["status"] == "ok" and r.json()["database"] is True


def test_cors_is_restricted_to_configured_origins(client):
    ok = client.get("/health", headers={"Origin": "http://localhost:3000"})
    assert ok.headers.get("access-control-allow-origin") == "http://localhost:3000"
    bad = client.get("/health", headers={"Origin": "https://evil.example"})
    assert "access-control-allow-origin" not in bad.headers


def test_job_lifecycle_and_test_data_isolation(client, fixture_site):
    job_id = crawl_fixture(client)
    prog = client.get(f"/api/jobs/{job_id}/progress").json()
    assert prog["status"] == "completed" and prog["percent"] == 100 and prog["successful_urls"] == 2
    assert client.get(f"/api/jobs/{job_id}").json()["is_test"] is True
    # sample leads are hidden unless explicitly requested
    assert client.get("/api/leads").json()["total"] == 0
    assert client.get("/api/leads", params={"include_test": True}).json()["total"] == 8
    assert client.get("/api/dashboard/stats").json()["unique_emails"] == 0
    stats = client.get("/api/dashboard/stats", params={"include_test": True}).json()
    assert stats["unique_emails"] == 8 and stats["websites_processed"] == 2 and stats["successful_crawls"] == 2 and stats["unique_businesses"] == 2
    assert stats["pending_jobs"] == 0 and stats["leads_today"] == 8 and stats["recent_jobs"][0]["id"] == job_id
    listing = client.get("/api/jobs").json()
    assert listing["total"] == 1 and listing["items"][0]["type"] == "website"


def test_lead_filters_search_sort_and_keyword_prospecting(client, fixture_site):
    crawl_fixture(client)
    g = lambda **p: client.get("/api/leads", params={"include_test": True, **p}).json()
    assert g(category="Packaging Distributor")["total"] == 7
    assert g(domain="greenleaf-foods.test")["total"] == 1
    assert g(state="Nevada")["total"] == 7 and g(city="Reno")["total"] == 7
    assert g(source_type="website")["total"] == 8 and g(source_type="upload")["total"] == 0
    assert g(has_website=True)["total"] == 8 and g(has_website=False)["total"] == 0
    assert g(q="orders")["total"] == 1 and g(q="Greenleaf")["total"] == 1
    assert g(verification_status="Domain MX Found")["total"] == 8
    assert g(role_based=True)["total"] >= 6
    assert g(date_from="2000-01-01", date_to="2999-12-31")["total"] == 8 and g(date_from="2999-01-01")["total"] == 0
    asc = [l["email"] for l in g(sort_by="email", sort_dir="asc")["items"]]
    assert asc == sorted(asc)
    page2 = g(page=2, page_size=3)
    assert len(page2["items"]) == 3 and page2["total"] == 8
    kw = g(keywords="custom packaging, shipping boxes", keyword_mode="all")
    assert kw["total"] == 7 and set(kw["items"][0]["matched_keywords"]) == {"custom packaging", "shipping boxes"}
    assert g(keywords="bakery")["total"] == 1  # only the Greenleaf page mentions a bakery
    opts = client.get("/api/leads/filters", params={"include_test": True}).json()
    assert "Packaging Distributor" in opts["categories"] and "Nevada" in opts["states"] and opts["has_test_data"]
    assert client.get("/api/leads", params={"date_from": "garbage"}).status_code == 422


def test_edit_tag_note_classify_and_delete_leads(client, fixture_site):
    crawl_fixture(client)
    items = client.get("/api/leads", params={"include_test": True, "page_size": 50}).json()["items"]
    lead = next(l for l in items if l["email"] == "sales@northwind-box.test")
    detail = client.get(f"/api/leads/{lead['id']}").json()
    assert detail["sources"] and detail["category_evidence"]["snippets"]
    r = client.patch(f"/api/leads/{lead['id']}", json={"business_name": "Northwind (edited)", "category": "Wholesale", "notes": "Called Monday", "city": "  "})
    assert r.status_code == 200 and r.json()["category_source"] == "manual" and r.json()["city"] is None and r.json()["notes"] == "Called Monday"
    client.post("/api/leads/classify", json={})
    assert client.get(f"/api/leads/{lead['id']}").json()["category"] == "Wholesale"  # manual choice survives re-classification
    ids = [l["id"] for l in items[:3]]
    assert client.post("/api/leads/bulk-tag", json={"ids": ids, "tags": ["hot", "q4"]}).json() == {"updated": 3}
    assert client.get("/api/leads", params={"include_test": True, "tag": "hot"}).json()["total"] == 3
    client.post("/api/leads/bulk-tag", json={"ids": ids, "tags": ["hot"], "mode": "remove"})
    assert client.get("/api/leads", params={"include_test": True, "tag": "hot"}).json()["total"] == 0
    assert client.post("/api/leads/bulk-delete", json={"ids": ids}).json() == {"deleted": 3}
    assert client.get("/api/leads", params={"include_test": True}).json()["total"] == 5
    assert client.patch("/api/leads/999999", json={}).status_code == 404
    assert client.patch(f"/api/leads/{items[5]['id']}", json={"verification_status": "Definitely Real"}).status_code == 422


def test_manual_lead_creation_and_validation_endpoint(client):
    assert client.post("/api/leads", json={"email": "Owner@Shop.com", "business_name": "Shop"}).status_code == 201
    assert client.post("/api/leads", json={"email": "owner@shop.com"}).status_code == 409
    assert client.post("/api/leads", json={"email": "nope"}).status_code == 422
    r = client.post("/api/leads/validate", json={}).json()
    assert r["checked"] == 1 and r["by_status"] == {"Domain MX Found": 1}


def test_excel_exports_all_scopes_and_errors(client, fixture_site):
    job_id = crawl_fixture(client)
    ids = [l["id"] for l in client.get("/api/leads", params={"include_test": True}).json()["items"]]

    def export(**body):
        r = client.post("/api/exports", json=body)
        if r.status_code != 201:
            return r
        d = client.get(r.json()["download_url"])
        assert d.status_code == 200 and "spreadsheetml" in d.headers["content-type"]
        return load_workbook(io.BytesIO(d.content)), r.json()

    assert client.post("/api/exports", json={"scope": "all"}).status_code == 400  # test data excluded -> nothing to export
    wb, rec = export(scope="all", include_test=True)
    assert wb["Leads"].max_row == 9 and rec["row_count"] == 8
    wb, rec = export(scope="selected", ids=ids[:2])
    assert rec["row_count"] == 2
    wb, rec = export(scope="filtered", include_test=True, filters={"domain": "greenleaf-foods.test", "include_test": True})
    assert rec["row_count"] == 1 and wb["Leads"]["B2"].value == "wholesale@greenleaf-foods.test"
    wb, rec = export(scope="all", include_test=True, format="unique_emails")
    assert wb.sheetnames == ["Unique Emails"]
    wb, rec = export(scope="all", include_test=True, format="by_domain")
    assert wb.sheetnames == ["Emails by Domain"]
    wb, rec = export(scope="job", job_id=job_id)
    assert rec["row_count"] == 8 and "Failed URLs" in wb.sheetnames
    assert client.post("/api/exports", json={"scope": "selected", "ids": []}).status_code == 400
    assert client.post("/api/exports", json={"scope": "job", "job_id": 999}).status_code == 404
    assert client.post("/api/exports", json={"scope": "nonsense"}).status_code == 422

    history = client.get("/api/exports").json()
    assert history["total"] == 6
    eid = history["items"][0]["id"]
    assert client.delete(f"/api/exports/{eid}").json() == {"deleted": True}
    assert client.get(f"/api/exports/{eid}/download").status_code == 404


def test_export_download_when_file_was_removed(client, fixture_site):
    crawl_fixture(client)
    rec = client.post("/api/exports", json={"scope": "all", "include_test": True}).json()
    from pathlib import Path
    from app.config import get_settings

    for f in Path(get_settings().export_dir).glob("*.xlsx"):
        f.unlink()
    r = client.get(rec["download_url"])
    assert r.status_code == 410 and "Create the export again" in r.json()["detail"]


# ------------------------------------------------------------------ uploads
def csv_bytes():
    return b"Company,Website,Email\nAcme Pack,acme-pack.com,buyer@acme-pack.com\nBeta,http://127.0.0.1:8099/,\nBuyer Inc,,hello@buyer-inc.com\nAcme Again,acme-pack.com,buyer@acme-pack.com\n"


def upload(client, name, data, ctype="text/csv"):
    return client.post("/api/uploads", files={"file": (name, data, ctype)})


def test_upload_preview_commit_emails_then_delete(client):
    r = upload(client, "contacts.csv", csv_bytes())
    assert r.status_code == 201
    up = r.json()
    sheet = up["sheets"][0]
    assert sheet["suggested_url_columns"] == [1] and sheet["suggested_email_columns"] == [2] and sheet["suggested_name_column"] == 0
    prev = client.post(f"/api/uploads/{up['upload_id']}/preview", json={}).json()
    assert prev["emails_total"] == 3 and prev["unique_emails"] == 2 and prev["duplicate_emails"] == 1 and prev["urls_total"] == 2
    acme = next(e for e in prev["emails"] if e["email"] == "buyer@acme-pack.com")
    assert acme["occurrences"] == 2 and {s["row"] for s in acme["sources"]} == {2, 5} and acme["sources"][0]["filename"] == "contacts.csv"
    assert client.get("/api/leads").json()["total"] == 0  # preview saves nothing
    done = client.post(f"/api/uploads/{up['upload_id']}/commit", json={"save_emails": True, "name": "My import"}).json()
    assert done["extract_job"]["created"] == 2 and done["extract_job"]["duplicates"] == 1 and done["file_deleted"] is True
    leads = client.get("/api/leads").json()
    assert leads["total"] == 2 and {l["source_filename"] for l in leads["items"]} == {"contacts.csv"} and all(l["source_type"] == "upload" for l in leads["items"])
    assert client.post(f"/api/uploads/{up['upload_id']}/preview", json={}).status_code == 404  # file removed after commit
    job = client.get(f"/api/jobs/{done['extract_job']['id']}").json()
    assert job["type"] == "upload" and job["input_filename"] == "contacts.csv" and job["emails_found"] == 3


def test_upload_commit_crawl_urls_with_manual_column_mapping(client, fixture_site):
    wb = Workbook()
    ws = wb.active
    ws.title = "Targets"
    ws.append(["Name", "Where", "Other"])
    ws.append(["Northwind", "http://127.0.0.1:8099/", "x"])
    ws.append(["Greenleaf", "http://localhost:8099/", "x"])
    buf = io.BytesIO()
    wb.save(buf)
    up = upload(client, "targets.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").json()
    assert up["sheets"][0]["name"] == "Targets"
    prev = client.post(f"/api/uploads/{up['upload_id']}/preview", json={"sheet": "Targets", "url_columns": [1], "name_column": 0}).json()
    assert prev["urls_total"] == 2 and prev["urls"][0]["business_name"] == "Northwind"
    client.post(f"/api/uploads/{up['upload_id']}/commit", json={"sheet": "Targets", "url_columns": [1], "name_column": 0, "save_emails": False, "crawl_urls": True, "delay": 0}).json()
    run_worker()
    leads = client.get("/api/leads", params={"include_test": True, "page_size": 50}).json()["items"]
    assert len(leads) == 8
    nw = next(l for l in leads if l["email"] == "sales@northwind-box.test")
    assert nw["business_name"] == "Northwind" and nw["source_filename"] == "targets.xlsx"  # the user's own name wins and the file is traceable
    detail = client.get(f"/api/leads/{nw['id']}").json()
    assert detail["sources"][0]["sheet_name"] == "Targets" and detail["sources"][0]["row_number"] == 2


def test_commit_requires_an_action(client):
    up = upload(client, "contacts.csv", csv_bytes()).json()
    assert client.post(f"/api/uploads/{up['upload_id']}/commit", json={"save_emails": False, "crawl_urls": False}).status_code == 400
    assert client.delete(f"/api/uploads/{up['upload_id']}").json() == {"deleted": True}
    assert client.delete(f"/api/uploads/{up['upload_id']}").status_code == 404


@pytest.mark.parametrize(
    "name,data,ctype,code,needle",
    [
        ("malware.exe", b"MZ\x90\x00", "application/octet-stream", 415, "not supported"),
        ("old.doc", b"\xd0\xcf\x11\xe0", "application/msword", 415, "docx"),
        ("fake.pdf", b"<html>not a pdf</html>", "application/pdf", 415, "not a PDF"),
        ("empty.csv", b"", "text/csv", 415, "empty"),
        ("wrong.csv", b"a,b\n1,2", "image/png", 415, "does not match"),
        ("../../evil.txt", b"hello", "text/plain", 201, ""),  # path components are stripped, never used on disk
    ],
)
def test_bad_uploads_are_rejected_clearly(client, name, data, ctype, code, needle):
    r = upload(client, name, data, ctype)
    assert r.status_code == code, r.text
    if code != 201:
        assert needle in r.json()["detail"]
    else:
        assert r.json()["filename"] == "evil.txt"


def test_corrupted_uploads_report_422_on_preview(client):
    up = upload(client, "broken.pdf", b"%PDF-1.4\ngarbage garbage garbage", "application/pdf").json()
    r = client.post(f"/api/uploads/{up['upload_id']}/preview", json={})
    assert r.status_code == 422 and "corrupted" in r.json()["detail"].lower()
    r = upload(client, "broken.xlsx", b"PK\x03\x04garbage", "application/octet-stream")
    assert r.status_code == 422


def test_oversized_upload_rejected(client, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "max_upload_mb", 0)
    r = upload(client, "big.txt", b"x" * 2048, "text/plain")
    assert r.status_code == 415 and "larger" in r.json()["detail"]


@pytest.mark.parametrize("bad_id", ["../../etc/passwd", "..%2F..%2Fetc%2Fpasswd", "a" * 31, "A" * 32, "0" * 32])
def test_upload_ids_cannot_traverse_paths(client, bad_id):
    assert client.post(f"/api/uploads/{bad_id}/preview", json={}).status_code in (404, 422)
    assert client.delete(f"/api/uploads/{bad_id}").status_code in (404, 405)


def test_scanned_pdf_upload_explains_ocr(client):
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    up = upload(client, "scan.pdf", doc.tobytes(), "application/pdf").json()
    prev = client.post(f"/api/uploads/{up['upload_id']}/preview", json={}).json()
    assert prev["needs_ocr"] and prev["emails_total"] == 0 and "Tesseract" in prev["warnings"][0]


def test_settings_roundtrip_and_custom_category_keywords_apply(client):
    s = client.get("/api/settings").json()
    assert s["settings"]["default_max_pages"] == 10 and s["server_limits"]["respect_robots_txt"] is True and "no-key" not in str(s)
    r = client.put("/api/settings", json={"default_max_pages": 25, "category_keywords": {"Eco Packaging": ["corrugated"]}})
    assert r.json()["settings"]["default_max_pages"] == 25
    assert client.put("/api/settings", json={"default_max_pages": 0}).status_code == 422


def test_invalid_job_requests(client):
    assert client.post("/api/jobs/website", json={"urls": []}).status_code == 422
    r = client.post("/api/jobs/website", json={"urls": ["bad url", "https://acme.com", "ACME.com"]}).json()
    assert (r["valid_urls"], r["invalid_urls"], r["duplicate_urls"]) == (1, 1, 1)
    assert client.get("/api/jobs/12345").status_code == 404
    assert client.post("/api/jobs/12345/cancel").status_code == 404
    job_id = r["job"]["id"]
    assert client.post(f"/api/jobs/{job_id}/pause").json()["status"] == "paused"
    assert client.post(f"/api/jobs/{job_id}/resume").json()["status"] == "queued"
    assert client.post(f"/api/jobs/{job_id}/cancel").json()["status"] == "cancelled"
    assert client.post(f"/api/jobs/{job_id}/cancel").status_code == 400
    assert client.get(f"/api/jobs/{job_id}/urls", params={"status": "failed"}).json()["total"] == 1
    assert client.delete(f"/api/jobs/{job_id}").status_code == 200
    assert client.get(f"/api/jobs/{job_id}").status_code == 404


def test_directory_job_endpoint(client, fixture_site):
    r = client.post("/api/jobs/directory", json={"urls": [fixture_site + "/directory/"], "delay": 0, "max_pages": 5, "validate_mx": False})
    assert r.status_code == 201
    run_worker()
    urls = client.get(f"/api/jobs/{r.json()['job']['id']}/urls").json()["items"]
    assert urls[0]["status"] == "success" and "4 businesses parsed" in urls[0]["note"] and "1 without a public email" in urls[0]["note"]
    leads = client.get("/api/leads", params={"include_test": True}).json()["items"]
    assert {l["email"] for l in leads} == {"contact@alphapack.test", "info@betabottling.test", "hello@gammafulfill.test"}
    assert all(l["source_type"] == "directory" for l in leads)

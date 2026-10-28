import io
from openpyxl import load_workbook

from app.models.job import Job, JobUrl
from app.models.lead import Lead
from app.services.excel_exporter import build_workbook, workbook_bytes
from app.services.lead_deduplicator import LeadCandidate, upsert_leads
from app.utils.security import sanitize_cell


def seed(db):
    cands = [
        LeadCandidate("sales@acme-pack.com", "Sales@Acme-Pack.com", "website", business_name="Acme Pack", website="https://acme-pack.com/", source_url="https://acme-pack.com/contact", state="Colorado", phone="+1 303 555 0100"),
        LeadCandidate("info@acme-pack.com", "info@acme-pack.com", "website", business_name="Acme Pack", website="https://acme-pack.com/", source_url="https://acme-pack.com/"),
        LeadCandidate("buyer@beta.org", "buyer@beta.org", "upload", business_name="=HYPERLINK(\"http://evil\",\"click\")", source_filename="leads.xlsx", sheet_name="Leads", row_number=4, extra={"Company": "Beta", "Notes": "+cmd"}),
        LeadCandidate("buyer@beta.org", "buyer@beta.org", "website", source_url="https://beta.org/contact"),  # duplicate
    ]
    upsert_leads(db, cands, job_id=None)
    db.commit()
    return db.query(Lead).order_by(Lead.id).all()


def wb_of(wb):
    return load_workbook(io.BytesIO(workbook_bytes(wb)))


def test_leads_workbook_structure(db):
    leads = seed(db)
    db.add(Job(name="J", type="website", status="completed", total_urls=2, processed_urls=2, failed_urls=1))
    db.flush()
    failed = [JobUrl(job_id=1, url="https://dead.com/", status="failed", error="HTTP 404")]
    wb = wb_of(build_workbook(leads, failed_urls=failed, jobs=db.query(Job).all()))
    assert wb.sheetnames == ["Leads", "Unique Emails", "Failed URLs", "Scraping Summary"]
    ws = wb["Leads"]
    headers = [c.value for c in ws[1]]
    for h in ("Business Name", "Email Address", "Original Email (as found)", "Website", "Contact Page URL", "Business Category", "City", "State",
              "Country", "Phone", "Source Type", "Source Filename", "Collection Date (UTC)", "Verification Status", "Notes", "All Sources", "Original Row Data"):
        assert h in headers, h
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and ws.max_row == 4
    row = {headers[i]: c.value for i, c in enumerate(ws[2])}
    assert row["Email Address"] and ws.column_dimensions["B"].width >= 12
    assert wb["Failed URLs"]["C2"].value == "HTTP 404"
    summary = {r[0].value: r[1].value for r in wb["Scraping Summary"].iter_rows(min_row=2)}
    assert summary["Leads in export"] == 3
    assert "NOT confirmed" in summary["Verification note"]


def test_original_and_normalized_values_both_exported(db):
    leads = seed(db)
    ws = wb_of(build_workbook(leads))["Leads"]
    h = [c.value for c in ws[1]]
    rows = [{h[i]: c.value for i, c in enumerate(r)} for r in ws.iter_rows(min_row=2)]
    sales = next(r for r in rows if r["Email Address"] == "sales@acme-pack.com")
    assert sales["Original Email (as found)"] == "Sales@Acme-Pack.com"
    beta = next(r for r in rows if r["Email Address"] == "buyer@beta.org")
    assert "leads.xlsx [Leads row 4]" in beta["All Sources"] and "https://beta.org/contact" in beta["All Sources"]  # both sightings kept
    assert '"Company": "Beta"' in beta["Original Row Data"]


def test_formula_injection_is_neutralised(db):
    leads = seed(db)
    ws = wb_of(build_workbook(leads))["Leads"]
    h = [c.value for c in ws[1]]
    beta = next(r for r in ws.iter_rows(min_row=2) if r[h.index("Email Address")].value == "buyer@beta.org")
    cell = beta[h.index("Business Name")]
    assert cell.value.startswith("'=HYPERLINK") and cell.data_type == "s"
    assert not any(c.data_type == "f" for r in ws.iter_rows() for c in r)
    assert sanitize_cell("@SUM(A1)") == "'@SUM(A1)" and sanitize_cell("-1+1") == "'-1+1" and sanitize_cell("+1 (303) 555-0100") == "+1 (303) 555-0100"
    assert sanitize_cell("ok\x00bad\x07") == "okbad"  # illegal XML chars removed instead of crashing openpyxl


def test_unique_emails_and_by_domain_scopes(db):
    leads = seed(db)
    wb = wb_of(build_workbook(leads, scope="unique_emails"))
    assert wb.sheetnames == ["Unique Emails"] and wb["Unique Emails"].max_row == 4
    wb = wb_of(build_workbook(leads, scope="by_domain"))
    ws = wb["Emails by Domain"]
    domains = [r[0].value for r in ws.iter_rows(min_row=2)]
    assert domains == sorted(domains) and domains.count("acme-pack.com") == 2
    assert [r[1].value for r in ws.iter_rows(min_row=2) if r[0].value == "acme-pack.com"] == [2, 2]


def test_empty_workbook_still_valid():
    wb = wb_of(build_workbook([]))
    assert wb["Leads"].max_row == 1  # header only
